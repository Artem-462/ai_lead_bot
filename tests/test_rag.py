import importlib.util
import io
import os
import sys
from unittest.mock import AsyncMock, MagicMock, patch

import numpy as np
import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

# Load real services/db.py
_db_path = os.path.join(PROJECT_ROOT, "services", "db.py")
_db_spec = importlib.util.spec_from_file_location("services.db.real", _db_path)
real_db = importlib.util.module_from_spec(_db_spec)
_db_spec.loader.exec_module(real_db)

# Load real services/rag.py
_rag_path = os.path.join(PROJECT_ROOT, "services", "rag.py")
_rag_spec = importlib.util.spec_from_file_location("services.rag", _rag_path)
rag = importlib.util.module_from_spec(_rag_spec)
sys.modules["services.rag"] = rag
_rag_spec.loader.exec_module(rag)

# Load real services/llm.py
_llm_path = os.path.join(PROJECT_ROOT, "services", "llm.py")
_llm_spec = importlib.util.spec_from_file_location("services.llm.real", _llm_path)
real_llm = importlib.util.module_from_spec(_llm_spec)
_llm_spec.loader.exec_module(real_llm)

import handlers.admin as admin_handler
from i18n import t
from tests.test_handlers import make_message


class TestPdfExtraction:
    def test_extract_pdf_pages_with_tables_and_text(self):
        page1 = MagicMock()
        table1 = MagicMock()
        table1.bbox = (10, 10, 100, 100)
        table1.extract.return_value = [
            [["Service", "Price"], ["Telegram Bot", "$300"], ["AI Assistant", "$500"]][0],
            [["Service", "Price"], ["Telegram Bot", "$300"], ["AI Assistant", "$500"]][1],
            [["Service", "Price"], ["Telegram Bot", "$300"], ["AI Assistant", "$500"]][2],
        ]
        page1.find_tables.return_value = [table1]
        filtered_page1 = MagicMock()
        filtered_page1.extract_text.return_value = "Overview of developer services."
        page1.filter.return_value = filtered_page1
        page1.extract_text.return_value = "Overview of developer services."

        page2 = MagicMock()
        page2.find_tables.return_value = []
        page2.extract_text.return_value = "Terms and conditions on page 2."

        fake_pdf = MagicMock()
        fake_pdf.pages = [page1, page2]
        fake_pdf.__enter__.return_value = fake_pdf
        fake_pdf.__exit__.return_value = None

        with patch("pdfplumber.open", return_value=fake_pdf):
            pages = rag.extract_pdf_pages(b"%PDF-mock")

        assert len(pages) == 2
        p1_num, p1_text = pages[0]
        assert p1_num == 1
        assert "Overview of developer services." in p1_text
        assert "[Table Row | Service: Telegram Bot | Price: $300]" in p1_text
        assert "[Table Row | Service: AI Assistant | Price: $500]" in p1_text

        p2_num, p2_text = pages[1]
        assert p2_num == 2
        assert "Terms and conditions on page 2." in p2_text

    def test_extract_pdf_pages_empty_scanned_pdf(self):
        page1 = MagicMock()
        page1.find_tables.return_value = []
        page1.extract_text.return_value = None

        fake_pdf = MagicMock()
        fake_pdf.pages = [page1]
        fake_pdf.__enter__.return_value = fake_pdf
        fake_pdf.__exit__.return_value = None

        with patch("pdfplumber.open", return_value=fake_pdf):
            pages = rag.extract_pdf_pages(b"%PDF-scanned")

        assert pages == []

    def test_format_table_row_multiline_and_carriage_return(self):
        headers = ["Col 1", "Col 2"]
        row = ["Line 1\r\nLine 2", "Text\nwith\rbreaks"]
        formatted = rag._format_table_row(headers, row)
        assert "\n" not in formatted
        assert "\r" not in formatted
        assert "[Table Row | Col 1: Line 1  Line 2 | Col 2: Text with breaks]" in formatted



class TestDocxExtraction:
    def test_docx_xml_page_break_and_tables(self):
        doc = Document()
        doc.add_paragraph("First page content")

        # Explicit page break
        p_break = doc.add_paragraph("Paragraph before break")
        p_break.add_run().add_break(WD_BREAK.PAGE)

        doc.add_paragraph("Second page content")

        # Add a table
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Feature"
        table.rows[0].cells[1].text = "Status"
        table.rows[1].cells[0].text = "RAG Engine"
        table.rows[1].cells[1].text = "Active"

        # Add cached Word rendered break
        p_rendered = doc.add_paragraph()
        run_rendered = p_rendered.add_run()
        rendered_break = OxmlElement("w:lastRenderedPageBreak")
        run_rendered._r.append(rendered_break)
        doc.add_paragraph("Third page content")

        buf = io.BytesIO()
        doc.save(buf)

        pages = rag.extract_docx_pages(buf.getvalue())
        assert len(pages) >= 2
        all_text = " ".join(p[1] for p in pages)
        assert "First page content" in all_text
        assert "Second page content" in all_text
        assert "[Table Row | Feature: RAG Engine | Status: Active]" in all_text

    def test_docx_fallback_pagination_1800_chars(self):
        doc = Document()
        for _ in range(30):
            doc.add_paragraph("A" * 90)

        buf = io.BytesIO()
        doc.save(buf)

        pages = rag.extract_docx_pages(buf.getvalue())
        assert len(pages) >= 2
        assert pages[0][0] == 1
        assert pages[1][0] == 2


class TestChunkPages:
    def test_chunking_preserves_table_rows_and_metadata(self):
        pages = [
            (
                1,
                (
                    "Introduction line.\n"
                    "[Table Row | Service: Web Bot | Price: $200]\n"
                    "[Table Row | Service: Voice Bot | Price: $400]\n"
                    "Conclusion line."
                ),
            )
        ]
        chunks = rag.chunk_pages(pages, chunk_size=300, overlap=50)
        assert len(chunks) >= 1
        for c in chunks:
            assert c["page_number"] == 1
            assert isinstance(c["chunk_index"], int)
            if "[Table Row |" in c["content"]:
                assert "Service:" in c["content"]
                assert "Price:" in c["content"]

    def test_chunking_wide_table_row_does_not_infinite_loop(self):
        wide_row = "[Table Row | " + "Detail: Data " * 80 + "]"
        pages = [
            (
                1,
                (
                    "Header line before table.\n"
                    "Another preparatory text.\n"
                    f"{wide_row}\n"
                    "Following paragraph text after the wide row."
                ),
            )
        ]
        chunks = rag.chunk_pages(pages, chunk_size=400, overlap=100)
        assert len(chunks) >= 2
        all_content = "\n".join(c["content"] for c in chunks)
        assert wide_row in all_content
        assert "Following paragraph text after the wide row." in all_content

    def test_chunk_pages_large_table_row_no_infinite_loop(self):
        short_paragraph = "A" * 100
        table_row = "[Table Row | " + "X" * (720 - len("[Table Row | ") - 1) + "]"
        assert len(short_paragraph) == 100
        assert len(table_row) == 720
        pages = [(1, f"{short_paragraph}\n{table_row}")]
        chunks = rag.chunk_pages(pages, chunk_size=800, overlap_size=150)
        assert len(chunks) == 2
        assert chunks[0]["content"] == short_paragraph
        assert chunks[1]["content"] == table_row


class TestEmbeddingCascade:
    def test_embedding_constants(self):
        assert rag.EMBEDDING_MODEL == "gemini-embedding-001"
        assert rag.FALLBACK_EMBEDDING_MODEL == "gemini-embedding-2"
        assert rag.EMBEDDING_MODELS == ("gemini-embedding-001", "gemini-embedding-2")

    @pytest.mark.asyncio
    async def test_get_embedding_fallback(self):
        sample_emb = MagicMock()
        sample_emb.values = [0.1, 0.2, 0.3]
        mock_response = MagicMock(embedding=sample_emb)

        called_models = []

        async def fake_embed_content(model, contents):
            called_models.append(model)
            if model == "gemini-embedding-001":
                raise rag.errors.APIError(429, "Rate limit")
            return mock_response

        with patch.object(rag.client.aio.models, "embed_content", side_effect=fake_embed_content):
            result = await rag.get_embedding("test text")

        assert called_models == ["gemini-embedding-001", "gemini-embedding-2"]
        assert np.allclose(result, np.array([0.1, 0.2, 0.3], dtype=np.float32))

    @pytest.mark.asyncio
    async def test_get_embeddings_fallback(self):
        sample_emb = MagicMock()
        sample_emb.values = [0.5, 0.6]
        mock_response = MagicMock(embeddings=[sample_emb])

        called_models = []

        async def fake_embed_content(model, contents):
            called_models.append(model)
            if model == "gemini-embedding-001":
                raise rag.errors.APIError(500, "Service unavailable")
            return mock_response

        with patch.object(rag.client.aio.models, "embed_content", side_effect=fake_embed_content):
            results = await rag.get_embeddings(["text1"])

        assert called_models == ["gemini-embedding-001", "gemini-embedding-2"]
        assert len(results) == 1
        assert np.allclose(results[0], np.array([0.5, 0.6], dtype=np.float32))



class TestDbRAGStore:
    @pytest.mark.asyncio
    async def test_upsert_versioning_and_delete(self, monkeypatch, tmp_db):
        monkeypatch.setattr(real_db, "DB_FILE", tmp_db)
        await real_db.init_db()

        sample_emb = np.array([0.1, 0.2, 0.3], dtype=np.float32)
        initial_chunks = [
            {"page_number": 1, "chunk_index": 0, "content": "Chunk 1", "embedding": sample_emb},
            {"page_number": 2, "chunk_index": 1, "content": "Chunk 2", "embedding": sample_emb},
        ]

        # Initial insert
        res1 = await real_db.upsert_document("catalog.pdf", "pdf", initial_chunks)
        assert res1["is_replaced"] is False
        assert res1["chunk_count"] == 2

        docs = await real_db.list_documents()
        assert len(docs) == 1
        assert docs[0]["filename"] == "catalog.pdf"
        assert docs[0]["chunk_count"] == 2

        all_chunks = await real_db.get_all_chunks_with_metadata()
        assert len(all_chunks) == 2

        # Re-uploading same filename replaces completely
        updated_chunks = [
            {"page_number": 1, "chunk_index": 0, "content": "New Chunk Only", "embedding": sample_emb}
        ]
        res2 = await real_db.upsert_document("catalog.pdf", "pdf", updated_chunks)
        assert res2["is_replaced"] is True
        assert res2["chunk_count"] == 1

        docs_after = await real_db.list_documents()
        assert len(docs_after) == 1
        assert docs_after[0]["chunk_count"] == 1

        chunks_after = await real_db.get_all_chunks_with_metadata()
        assert len(chunks_after) == 1
        assert chunks_after[0]["content"] == "New Chunk Only"

        # Deletion
        del_res = await real_db.delete_document_by_name("catalog.pdf")
        assert del_res is True
        assert await real_db.list_documents() == []
        assert await real_db.get_all_chunks_with_metadata() == []

        # Deleting nonexistent
        del_nonexistent = await real_db.delete_document_by_name("catalog.pdf")
        assert del_nonexistent is False

    @pytest.mark.asyncio
    async def test_upsert_and_delete_case_insensitive_collation(self, monkeypatch, tmp_db):
        monkeypatch.setattr(real_db, "DB_FILE", tmp_db)
        await real_db.init_db()

        sample_emb = np.array([0.1, 0.2, 0.3], dtype=np.float32)
        initial_chunks = [
            {"page_number": 1, "chunk_index": 0, "content": "Version 1", "embedding": sample_emb}
        ]
        res1 = await real_db.upsert_document("Report.pdf", "pdf", initial_chunks)
        assert res1["is_replaced"] is False

        updated_chunks = [
            {"page_number": 1, "chunk_index": 0, "content": "Version 2", "embedding": sample_emb}
        ]
        res2 = await real_db.upsert_document("report.PDF", "pdf", updated_chunks)
        assert res2["is_replaced"] is True

        docs = await real_db.list_documents()
        assert len(docs) == 1

        deleted = await real_db.delete_document_by_name("REPORT.pdf")
        assert deleted is True
        assert await real_db.list_documents() == []



class TestVectorSearch:
    @pytest.mark.asyncio
    async def test_search_knowledge_base_cosine_threshold(self):
        query_vector = np.array([1.0, 0.0, 0.0], dtype=np.float32)
        match_vector = np.array([0.9, 0.1, 0.0], dtype=np.float32)  # High similarity > 0.9
        low_vector = np.array([0.1, 0.9, 0.0], dtype=np.float32)    # Low similarity ~ 0.1

        mock_chunks = [
            {
                "content": "Relevant content",
                "filename": "guide.pdf",
                "page_number": 3,
                "chunk_index": 0,
                "embedding": match_vector,
            },
            {
                "content": "Irrelevant content",
                "filename": "other.docx",
                "page_number": 1,
                "chunk_index": 0,
                "embedding": low_vector,
            },
        ]

        with patch("services.rag.get_all_chunks_with_metadata", new_callable=AsyncMock, return_value=mock_chunks), \
             patch("services.rag.get_embedding", new_callable=AsyncMock, return_value=query_vector):
            results = await rag.search_knowledge_base("how to configure", top_k=2, min_score=0.52)

        assert len(results) == 1
        assert results[0]["filename"] == "guide.pdf"
        assert results[0]["page_number"] == 3
        assert results[0]["score"] > 0.52



class TestStrictRagMode:
    @pytest.mark.asyncio
    async def test_rag_refusal_on_empty_search_results(self):
        with patch.object(real_llm, "has_documents", new_callable=AsyncMock, return_value=True), \
             patch("services.rag.search_knowledge_base", new_callable=AsyncMock, return_value=[]):
            res_uk = await real_llm.get_ai_answer(123, "What is the secret?", "uk")
            res_en = await real_llm.get_ai_answer(123, "What is the secret?", "en")

        assert res_uk == t("rag_no_answer", "uk")
        assert res_en == t("rag_no_answer", "en")

    @pytest.mark.asyncio
    async def test_rag_refusal_when_model_outputs_no_rag_answer(self):
        chunks = [{"content": "About prices", "filename": "price.pdf", "page_number": 1}]
        with patch.object(real_llm, "has_documents", new_callable=AsyncMock, return_value=True), \
             patch("services.rag.search_knowledge_base", new_callable=AsyncMock, return_value=chunks), \
             patch.object(real_llm, "_call_model", new_callable=AsyncMock, return_value="NO_RAG_ANSWER"):
            res = await real_llm.get_ai_answer(123, "Tell me a joke", "uk")

        assert res == t("rag_no_answer", "uk")

    @pytest.mark.asyncio
    async def test_rag_successful_answer_with_deterministic_footer(self):
        chunks = [
            {"content": "Bot cost is $300", "filename": "price.pdf", "page_number": 2},
            {"content": "Support is $50/mo", "filename": "price.pdf", "page_number": 3},
        ]
        with patch.object(real_llm, "has_documents", new_callable=AsyncMock, return_value=True), \
             patch("services.rag.search_knowledge_base", new_callable=AsyncMock, return_value=chunks), \
             patch.object(real_llm, "_call_model", new_callable=AsyncMock, return_value="Ціна бота від $300."):
            res_uk = await real_llm.get_ai_answer(123, "Яка ціна бота?", "uk")

        assert "Ціна бота від $300." in res_uk
        assert "📄 Джерело: price.pdf (стор. 2, 3)" in res_uk

        with patch.object(real_llm, "has_documents", new_callable=AsyncMock, return_value=True), \
             patch("services.rag.search_knowledge_base", new_callable=AsyncMock, return_value=chunks), \
             patch.object(real_llm, "_call_model", new_callable=AsyncMock, return_value="Bot price starts at $300."):
            res_en = await real_llm.get_ai_answer(123, "What is the bot price?", "en")

        assert "Bot price starts at $300." in res_en
        assert "📄 Source: price.pdf (p. 2, 3)" in res_en



class TestAdminRagHandlers:
    @pytest.mark.asyncio
    async def test_cmd_docs_non_admin_forbidden(self):
        msg = make_message(text="/docs", user_id=111)
        with patch("config.ADMIN_ID", 999):
            await admin_handler.cmd_docs(msg)
        assert msg.answer.await_count == 1
        assert msg.answer.call_args[0][0] == t("stats_no_access", "uk")

    @pytest.mark.asyncio
    async def test_cmd_docs_empty_and_populated(self):
        msg = make_message(text="/docs", user_id=999)
        with patch("config.ADMIN_ID", 999), \
             patch("handlers.admin.list_documents", new_callable=AsyncMock, return_value=[]):
            await admin_handler.cmd_docs(msg)
        assert msg.answer.call_args[0][0] == t("doc_list_empty", "uk")

        msg2 = make_message(text="/docs", user_id=999)
        docs = [{"filename": "info.pdf", "file_type": "pdf", "chunk_count": 5, "uploaded_at": "2026-10-20"}]
        with patch("config.ADMIN_ID", 999), \
             patch("handlers.admin.list_documents", new_callable=AsyncMock, return_value=docs):
            await admin_handler.cmd_docs(msg2)
        assert "info.pdf" in msg2.answer.call_args[0][0]

    @pytest.mark.asyncio
    async def test_cmd_del_doc(self):
        msg = make_message(text="/del_doc guide.pdf", user_id=999)
        with patch("config.ADMIN_ID", 999), \
             patch("handlers.admin.delete_document_by_name", new_callable=AsyncMock, return_value=True):
            await admin_handler.cmd_del_doc(msg)
        assert t("doc_deleted", "uk", filename="guide.pdf") in msg.answer.call_args[0][0]

        msg_fail = make_message(text="/del_doc nonexistent.pdf", user_id=999)
        with patch("config.ADMIN_ID", 999), \
             patch("handlers.admin.delete_document_by_name", new_callable=AsyncMock, return_value=False):
            await admin_handler.cmd_del_doc(msg_fail)
        assert t("doc_not_found", "uk", filename="nonexistent.pdf") in msg_fail.answer.call_args[0][0]

    @pytest.mark.asyncio
    async def test_handle_document_upload(self):
        # Non-admin
        msg_non_admin = make_message(user_id=111)
        msg_non_admin.document = MagicMock()
        msg_non_admin.document.file_name = "test.pdf"
        with patch("config.ADMIN_ID", 999):
            await admin_handler.handle_document_upload(msg_non_admin)
        assert msg_non_admin.answer.call_args[0][0] == t("non_text", "uk")

        # Unsupported format
        msg_unsupported = make_message(user_id=999)
        msg_unsupported.document = MagicMock()
        msg_unsupported.document.file_name = "script.exe"
        with patch("config.ADMIN_ID", 999):
            await admin_handler.handle_document_upload(msg_unsupported)
        assert msg_unsupported.answer.call_args[0][0] == t("doc_unsupported", "uk")

        # Valid PDF upload (new)
        msg_pdf = make_message(user_id=999)
        msg_pdf.document = MagicMock()
        msg_pdf.document.file_name = "pricing.pdf"
        wait_msg = AsyncMock()
        msg_pdf.answer = AsyncMock(return_value=wait_msg)
        msg_pdf.bot.download = AsyncMock(return_value=io.BytesIO(b"dummy pdf bytes"))

        with patch("config.ADMIN_ID", 999), \
             patch("handlers.admin.index_document", new_callable=AsyncMock, return_value={"is_replaced": False, "chunk_count": 4}):
            await admin_handler.handle_document_upload(msg_pdf)

        wait_msg.edit_text.assert_awaited_once()
        assert t("doc_indexed_new", "uk", filename="pricing.pdf", chunks=4) in wait_msg.edit_text.call_args[0][0]

    @pytest.mark.asyncio
    async def test_admin_handlers_html_escape_filenames(self):
        # 1. cmd_docs
        msg_docs = make_message(text="/docs", user_id=999)
        docs = [{"filename": "<doc>&test.pdf", "file_type": "pdf", "chunk_count": 3, "uploaded_at": "2026-10-21"}]
        with patch("config.ADMIN_ID", 999), \
             patch("handlers.admin.list_documents", new_callable=AsyncMock, return_value=docs):
            await admin_handler.cmd_docs(msg_docs)
        assert "&lt;doc&gt;&amp;test.pdf" in msg_docs.answer.call_args[0][0]
        assert "<doc>" not in msg_docs.answer.call_args[0][0]

        # 2. cmd_del_doc
        msg_del = make_message(text="/del_doc <del>&doc.pdf", user_id=999)
        with patch("config.ADMIN_ID", 999), \
             patch("handlers.admin.delete_document_by_name", new_callable=AsyncMock, return_value=True):
            await admin_handler.cmd_del_doc(msg_del)
        assert "&lt;del&gt;&amp;doc.pdf" in msg_del.answer.call_args[0][0]
        assert "<del>" not in msg_del.answer.call_args[0][0]

        # 3. handle_document_upload
        msg_upload = make_message(user_id=999)
        msg_upload.document = MagicMock()
        msg_upload.document.file_name = "<upload>&file.pdf"
        wait_msg = AsyncMock()
        msg_upload.answer = AsyncMock(return_value=wait_msg)
        msg_upload.bot.download = AsyncMock(return_value=io.BytesIO(b"dummy pdf bytes"))

        with patch("config.ADMIN_ID", 999), \
             patch("handlers.admin.index_document", new_callable=AsyncMock, return_value={"is_replaced": False, "chunk_count": 2}):
            await admin_handler.handle_document_upload(msg_upload)

        indexing_text = msg_upload.answer.call_args[0][0]
        assert "&lt;upload&gt;&amp;file.pdf" in indexing_text
        assert "<upload>" not in indexing_text
        result_text = wait_msg.edit_text.call_args[0][0]
        assert "&lt;upload&gt;&amp;file.pdf" in result_text
        assert "<upload>" not in result_text

    @pytest.mark.asyncio
    async def test_handle_document_upload_file_too_large(self):
        msg = make_message(user_id=999)
        msg.document = MagicMock()
        msg.document.file_name = "huge.pdf"
        msg.document.file_size = 25 * 1024 * 1024

        with patch("config.ADMIN_ID", 999):
            await admin_handler.handle_document_upload(msg)

        msg.answer.assert_awaited_once()
        assert t("doc_too_large", "uk") in msg.answer.call_args[0][0]
        msg.bot.download.assert_not_called()

    @pytest.mark.asyncio
    async def test_handle_document_upload_empty_document(self):
        msg = make_message(user_id=999)
        msg.document = MagicMock()
        msg.document.file_name = "empty.pdf"
        msg.document.file_size = 1024
        wait_msg = AsyncMock()
        msg.answer = AsyncMock(return_value=wait_msg)
        msg.bot.download = AsyncMock(return_value=io.BytesIO(b"empty bytes"))

        with patch("config.ADMIN_ID", 999), \
             patch("handlers.admin.index_document", new_callable=AsyncMock, side_effect=ValueError("No text extracted from document: empty.pdf")):
            await admin_handler.handle_document_upload(msg)

        wait_msg.edit_text.assert_awaited_once()
        assert t("doc_empty", "uk", filename="empty.pdf") in wait_msg.edit_text.call_args[0][0]
